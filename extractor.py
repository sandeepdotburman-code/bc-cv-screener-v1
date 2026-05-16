# extractor.py
# Takes a CV file (PDF or DOCX) and returns its contents as plain text.
# Supports text-layer PDFs, image-based PDFs (via OCR), DOCX, and flags images.
# Called by screener.py — one file at a time.

import pdfplumber
from docx import Document
import pytesseract
from pdf2image import convert_from_path
import os

# ── Path configuration ───────────────────────────────────────────────────────
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
POPPLER_PATH = r'C:\Program Files\poppler-26.02.0\Library\bin'

def extract_text(file_path: str) -> dict:
    """
    Extract text from a PDF or DOCX file.

    Returns a dict with:
      - 'text': the extracted text string (empty string if extraction failed)
      - 'status': 'ok', 'unreadable', or 'unsupported'
      - 'reason': human-readable explanation if status is not 'ok'
      - 'format': file format detected (PDF, DOCX, JPEG, Other)
    """

    ext = os.path.splitext(file_path)[1].lower()

    # ── JPEG and other image formats ─────────────────────────────────────────
    if ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff']:
        try:
            from PIL import Image
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
            if text.strip():
                return {
                    'text': text,
                    'status': 'ok',
                    'reason': '',
                    'format': 'JPEG'
                }
        except Exception:
            pass
        return {
            'text': '',
            'status': 'unreadable',
            'reason': f'{os.path.basename(file_path)} — JPEG — no text extracted. Navdeep to process manually.',
            'format': 'JPEG'
        }

    # ── DOCX ─────────────────────────────────────────────────────────────────
    if ext == '.docx':
        try:
            doc = Document(file_path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            full_text = '\n'.join(paragraphs)
            if not full_text.strip():
                return {
                    'text': '',
                    'status': 'unreadable',
                    'reason': f'{os.path.basename(file_path)} — DOCX — no text extracted. Navdeep to process manually.',
                    'format': 'DOCX'
                }
            return {
                'text': full_text,
                'status': 'ok',
                'reason': '',
                'format': 'DOCX'
            }
        except Exception as e:
            return {
                'text': '',
                'status': 'unreadable',
                'reason': f'{os.path.basename(file_path)} — DOCX — could not be read ({str(e)}). Navdeep to process manually.',
                'format': 'DOCX'
            }

    # ── PDF ──────────────────────────────────────────────────────────────────
    if ext == '.pdf':
        # First try direct text extraction (fast — works on text-layer PDFs)
        try:
            full_text_parts = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        full_text_parts.append(page_text)
            full_text = '\n'.join(full_text_parts)

            if full_text.strip():
                return {
                    'text': full_text,
                    'status': 'ok',
                    'reason': '',
                    'format': 'PDF'
                }
        except Exception:
            pass

        # No text layer found — fall back to OCR
        try:
            images = convert_from_path(file_path, poppler_path=POPPLER_PATH)
            ocr_parts = []
            for image in images:
                page_text = pytesseract.image_to_string(image)
                if page_text.strip():
                    ocr_parts.append(page_text)
            full_text = '\n'.join(ocr_parts)

            if full_text.strip():
                return {
                    'text': full_text,
                    'status': 'ok',
                    'reason': '',
                    'format': 'PDF'
                }
            else:
                return {
                    'text': '',
                    'status': 'unreadable',
                    'reason': f'{os.path.basename(file_path)} — PDF — no text extracted after OCR attempt. Navdeep to process manually.',
                    'format': 'PDF'
                }
        except Exception as e:
            return {
                'text': '',
                'status': 'unreadable',
                'reason': f'{os.path.basename(file_path)} — PDF — OCR failed ({str(e)}). Navdeep to process manually.',
                'format': 'PDF'
            }

    # ── Unsupported format ───────────────────────────────────────────────────
    return {
        'text': '',
        'status': 'unsupported',
        'reason': f'{os.path.basename(file_path)} — {ext.upper().strip(".")} — unsupported format. Navdeep to process manually.',
        'format': 'Other'
    }